
import pandas as pd
from pathlib import Path
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# AUTOMATIZADOR DE RELATÓRIOS EXCEL
PASTA_PROJETO = Path(__file__).resolve().parent

base_dados = None
arquivo_origem = None

# 1. IMPORTAR PLANILHA
def importar_planilha():
    global base_dados, arquivo_origem

    print("\n--- IMPORTAR PLANILHA EXCEL ---")

    caminho = input(
        "Digite o caminho do arquivo Excel: "
    ).strip().strip('"').strip("'")

    arquivo = Path(caminho).expanduser()

    if not arquivo.is_file():
        print("\nArquivo não encontrado.")
        return

    if arquivo.suffix.lower() not in [".xlsx", ".xls"]:
        print("\nSelecione um arquivo Excel válido.")
        return

    try:
        planilha = pd.ExcelFile(arquivo)

        print("\nAbas disponíveis:")

        for indice, nome in enumerate(
            planilha.sheet_names, start=1
        ):
            print(f"[{indice}] {nome}")

        escolha = input(
            "\nSelecione o número da aba: "
        ).strip()

        if not escolha.isdigit():
            print("Opção inválida.")
            return

        indice = int(escolha) - 1

        if indice not in range(len(planilha.sheet_names)):
            print("Aba não encontrada.")
            return

        dados = pd.read_excel(
            planilha,
            sheet_name=indice
        )

        if dados.empty:
            print("\nA aba selecionada está vazia.")
            return

        base_dados = dados
        arquivo_origem = arquivo

        print("\nPlanilha importada com sucesso!")
        print(f"Arquivo: {arquivo.name}")
        print(f"Aba: {planilha.sheet_names[indice]}")
        print(f"Registros: {len(base_dados)}")

    except Exception as erro:
        print(f"\nErro ao importar: {erro}")

# 2. VERIFICAR BASE IMPORTADA
def verificar_base():

    if base_dados is None:
        print("\nImporte uma planilha primeiro.")
        return False

    return True

# 3. VISUALIZAR DADOS
def visualizar_dados():

    if not verificar_base():
        return

    print("\n--- VISUALIZAÇÃO DA BASE ---")

    print(
        base_dados.head(15).to_string(
            index=False
        )
    )

    print(
        f"\nTotal de registros: {len(base_dados)}"
    )

# 4. CONSULTAR INFORMAÇÕES DA BASE
def consultar_base():

    if not verificar_base():
        return

    print("\n--- INFORMAÇÕES DA BASE ---")

    print(f"Linhas: {base_dados.shape[0]}")
    print(f"Colunas: {base_dados.shape[1]}")

    print("\nColunas identificadas:")

    for coluna in base_dados.columns:
        print(f"- {coluna}")

    print("\nTipos de dados:")

    print(base_dados.dtypes.to_string())

    print("\nValores ausentes por coluna:")

    print(
        base_dados.isnull().sum().to_string()
    )

    duplicados = base_dados.duplicated().sum()

    print(
        f"\nRegistros duplicados: {duplicados}"
    )
# 5. ANALISAR INDICADORES
def calcular_indicadores():

    if not verificar_base():
        return None

    colunas_numericas = base_dados.select_dtypes(
        include="number"
    )

    if colunas_numericas.empty:
        print("\nNenhuma coluna numérica encontrada.")
        return None

    indicadores = pd.DataFrame({
        "Coluna": colunas_numericas.columns,
        "Total": colunas_numericas.sum().values,
        "Média": colunas_numericas.mean().values,
        "Mínimo": colunas_numericas.min().values,
        "Máximo": colunas_numericas.max().values,
        "Valores Válidos": colunas_numericas.count().values
    })

    return indicadores


def analisar_indicadores():

    indicadores = calcular_indicadores()

    if indicadores is None:
        return

    print("\n--- INDICADORES DA PLANILHA ---")

    print(
        indicadores.to_string(
            index=False
        )
    )

    print(
        "\nObservação: os cálculos são realizados "
        "sobre todas as colunas numéricas."
    )

# 6. FORMATAR ARQUIVO EXCEL
def formatar_aba(aba):

    cor_cabecalho = "17365D"
    cor_texto = "FFFFFF"
    cor_linhas = "EAF2F8"

    aba.freeze_panes = "A2"

    aba.auto_filter.ref = aba.dimensions

    for celula in aba[1]:

        celula.font = Font(
            bold=True,
            color=cor_texto
        )

        celula.fill = PatternFill(
            fill_type="solid",
            fgColor=cor_cabecalho
        )

        celula.alignment = Alignment(
            horizontal="center"
        )

    for linha in range(2, aba.max_row + 1):

        if linha % 2 == 0:

            for celula in aba[linha]:

                celula.fill = PatternFill(
                    fill_type="solid",
                    fgColor=cor_linhas
                )

    for coluna in aba.columns:

        letra = get_column_letter(
            coluna[0].column
        )

        maior = max(
            len(str(celula.value or ""))
            for celula in coluna
        )

        aba.column_dimensions[letra].width = min(
            max(maior + 3, 14),
            40
        )

# 7. GERAR RELATÓRIO AUTOMÁTICO
def gerar_relatorio():

    if not verificar_base():
        return

    indicadores = calcular_indicadores()

    resumo = pd.DataFrame({
        "Indicador": [
            "Total de registros",
            "Total de colunas",
            "Registros duplicados",
            "Valores ausentes",
            "Data da análise"
        ],
        "Resultado": [
            len(base_dados),
            len(base_dados.columns),
            int(base_dados.duplicated().sum()),
            int(base_dados.isnull().sum().sum()),
            datetime.now().strftime(
                "%d/%m/%Y %H:%M:%S"
            )
        ]
    })

    qualidade = pd.DataFrame({
        "Coluna": [str(c) for c in base_dados.columns],
        "Tipo de dado": [
            str(tipo) for tipo in base_dados.dtypes
        ],
        "Valores ausentes": (
            base_dados.isnull().sum().values
        ),
        "Valores preenchidos": (
            base_dados.notnull().sum().values
        )
    })

    nome_arquivo = (
        "relatorio_"
        + datetime.now().strftime("%Y%m%d_%H%M%S")
        + ".xlsx"
    )

    destino = PASTA_PROJETO / nome_arquivo

    try:

        with pd.ExcelWriter(
            destino,
            engine="openpyxl",
            datetime_format="DD/MM/YYYY HH:MM:SS"
        ) as escritor:

            resumo.to_excel(
                escritor,
                sheet_name="Resumo Geral",
                index=False
            )

            qualidade.to_excel(
                escritor,
                sheet_name="Qualidade dos Dados",
                index=False
            )

            if indicadores is not None:

                indicadores.to_excel(
                    escritor,
                    sheet_name="Indicadores",
                    index=False
                )

            base_dados.to_excel(
                escritor,
                sheet_name="Base Importada",
                index=False
            )

            for aba in escritor.book.worksheets:
                formatar_aba(aba)

        print("\nRelatório gerado com sucesso!")
        print(f"Arquivo: {destino}")
        print("\nO arquivo original não foi alterado.")

    except Exception as erro:

        print(
            f"\nErro ao gerar relatório: {erro}"
        )

# 8. MENU PRINCIPAL
def menu():

    while True:

        print("\n" + "=" * 48)
        print("     AUTOMATIZADOR DE RELATÓRIOS EXCEL")
        print("=" * 48)

        print("[1] Importar planilha Excel")
        print("[2] Visualizar dados importados")
        print("[3] Consultar informações da base")
        print("[4] Analisar indicadores")
        print("[5] Gerar relatório Excel automático")
        print("[0] Sair")

        opcao = input(
            "\nEscolha uma opção: "
        ).strip()

        if opcao == "1":
            importar_planilha()

        elif opcao == "2":
            visualizar_dados()

        elif opcao == "3":
            consultar_base()

        elif opcao == "4":
            analisar_indicadores()

        elif opcao == "5":
            gerar_relatorio()

        elif opcao == "0":
            print("\nPrograma encerrado.")
            break

        else:
            print("\nOpção inválida.")

# 9. EXECUÇÃO
if __name__ == "__main__":

    try:
        menu()

    except KeyboardInterrupt:
        print("\nPrograma interrompido.")
